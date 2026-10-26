from docx import Document
from docx.shared import Pt
import os
import time
from config import REPORT_OUTPUT_DIR

def generate_word_document(title: str, content: str, filename_prefix: str = "report") -> str:
    """
    生成 Word 文档
    返回文件路径
    """
    os.makedirs(REPORT_OUTPUT_DIR, exist_ok=True)
    
    doc = Document()
    doc.add_heading(title, level=1)
    
    # 按段落分割内容（Markdown 简易转换）
    paragraphs = content.split('\n\n')
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if para.startswith('# '):
            doc.add_heading(para[2:], level=2)
        elif para.startswith('## '):
            doc.add_heading(para[3:], level=3)
        elif para.startswith('### '):
            doc.add_heading(para[4:], level=4)
        elif para.startswith('- '):
            for line in para.split('\n'):
                if line.startswith('- '):
                    doc.add_paragraph(line[2:], style='List Bullet')
        elif para.startswith('1. ') or para.startswith('2. ') or para.startswith('3. '):
            doc.add_paragraph(para[3:], style='List Number')
        else:
            doc.add_paragraph(para)
    
    filename = f"{filename_prefix}_{int(time.time())}.docx"
    file_path = os.path.join(REPORT_OUTPUT_DIR, filename)
    doc.save(file_path)
    return file_path
